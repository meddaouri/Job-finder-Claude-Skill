#!/usr/bin/env python3
"""
Job Finder — build an ATS-friendly CV (.docx) from a JSON profile.

Usage:
    python3 build_cv_docx.py cv.json CV_Firstname_Lastname.docx

Requires: python-docx  (pip install python-docx)

JSON format (only "name" is required; omit what you don't have — never invent):
{
  "lang": "en",                 # en | fr | es | ar | de | pt  (section headings); any other → en, or give "labels"
  "format": "generic",          # generic | fr | us | uk | ca | de | europass | gulf  (page size + conventions)
  "rtl": false,                 # true for Arabic/Hebrew/Persian/Urdu CVs (auto true when lang = "ar")
  "name": "Jane Doe",
  "title": "Registered Nurse",
  "contact": {"city": "Lyon", "country": "France", "phone": "+33 …", "email": "…", "links": ["linkedin.com/in/…"]},
  "photo": "",                  # optional path to a photo the USER provided; used only if non-empty
  "summary": "2–3 lines",
  "experience": [{"role": "", "employer": "", "city": "", "start": "09/2021", "end": "Present", "bullets": ["…"]}],
  "education":  [{"degree": "", "school": "", "city": "", "year": "", "details": ""}],
  "skills":     [{"group": "Technical", "items": ["…", "…"]}],
  "languages":  [{"name": "French", "level": "Native"}],
  "certifications": ["…"],
  "projects":   ["…"],
  "extras":     ["…"],          # volunteering, interests…
  "education_first": null,      # true/false; null = auto (true when there is no experience)
  "labels": {}                  # optional overrides for headings, e.g. {"experience": "Berufserfahrung"}
}
"""
import json
import sys

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor, Inches
except ImportError:
    sys.exit("python-docx is missing: pip install python-docx")

LABELS = {
    "en": {"summary": "Profile", "experience": "Professional Experience", "education": "Education", "skills": "Skills",
           "languages": "Languages", "certifications": "Certifications", "projects": "Projects", "extras": "Additional Information"},
    "fr": {"summary": "Profil", "experience": "Expérience professionnelle", "education": "Formation", "skills": "Compétences",
           "languages": "Langues", "certifications": "Certifications", "projects": "Projets", "extras": "Informations complémentaires"},
    "es": {"summary": "Perfil", "experience": "Experiencia profesional", "education": "Formación", "skills": "Habilidades",
           "languages": "Idiomas", "certifications": "Certificaciones", "projects": "Proyectos", "extras": "Información adicional"},
    "ar": {"summary": "الملف الشخصي", "experience": "الخبرة المهنية", "education": "التعليم", "skills": "المهارات",
           "languages": "اللغات", "certifications": "الشهادات المهنية", "projects": "المشاريع", "extras": "معلومات إضافية"},
    "de": {"summary": "Profil", "experience": "Berufserfahrung", "education": "Ausbildung", "skills": "Kenntnisse",
           "languages": "Sprachen", "certifications": "Zertifikate", "projects": "Projekte", "extras": "Sonstiges"},
    "pt": {"summary": "Perfil", "experience": "Experiência profissional", "education": "Formação", "skills": "Competências",
           "languages": "Idiomas", "certifications": "Certificações", "projects": "Projetos", "extras": "Informações adicionais"},
}
LETTER_FORMATS = {"us", "ca"}
ACCENT = RGBColor(0x1E, 0x23, 0x30)
FONT = "Calibri"


def set_rtl(paragraph):
    """Mark a paragraph right-to-left. Bidi paragraphs start on the right by default, so alignment is left
    untouched (setting 'right' would flip them to the left). Complex-script runs need their own bold/size."""
    pPr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    pPr.append(bidi)
    for run in paragraph.runs:
        rPr = run._r.get_or_add_rPr()
        if run.bold:
            rPr.append(OxmlElement("w:bCs"))
        if run.font.size:
            szcs = OxmlElement("w:szCs")
            szcs.set(qn("w:val"), str(int(run.font.size.pt * 2)))
            rPr.append(szcs)
        rtl = OxmlElement("w:rtl")
        rtl.set(qn("w:val"), "1")
        rPr.append(rtl)


def add_bottom_border(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "9AA0AE")):
        bottom.set(qn(k), v)
    pbdr.append(bottom)
    pPr.append(pbdr)


class CV:
    def __init__(self, data):
        self.d = data
        lang = (data.get("lang") or "en").lower()[:2]
        self.labels = dict(LABELS.get(lang, LABELS["en"]))
        self.labels.update(data.get("labels") or {})
        rtl = data.get("rtl")
        self.rtl = (lang in ("ar", "he", "fa", "ur")) if rtl is None else bool(rtl)
        self.doc = Document()
        self._setup()

    def _setup(self):
        sec = self.doc.sections[0]
        if (self.d.get("format") or "").lower() in LETTER_FORMATS:
            sec.page_width, sec.page_height = Inches(8.5), Inches(11)
        else:
            sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        for side in ("left_margin", "right_margin"):
            setattr(sec, side, Cm(1.8))
        sec.top_margin = sec.bottom_margin = Cm(1.5)
        st = self.doc.styles["Normal"]
        st.font.name = FONT
        st.font.size = Pt(10.5)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        st.element.rPr.rFonts.set(qn("w:cs"), "Arial")
        st.paragraph_format.space_after = Pt(2)

    def p(self, text="", size=None, bold=False, italic=False, color=None, align=None, after=2, before=0):
        para = self.doc.add_paragraph()
        if text:
            r = para.add_run(text)
            r.bold, r.italic = bold, italic
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = color
        pf = para.paragraph_format
        pf.space_after, pf.space_before = Pt(after), Pt(before)
        if align is not None:
            para.alignment = align
        if self.rtl:
            set_rtl(para)
        return para

    def heading(self, key):
        para = self.p(self.labels[key].upper() if not self.rtl else self.labels[key], size=11.5, bold=True,
                      color=ACCENT, before=10, after=4)
        add_bottom_border(para)

    def bullet(self, text):
        para = self.doc.add_paragraph(style="List Bullet")
        para.add_run(text)
        para.paragraph_format.space_after = Pt(1)
        if self.rtl:
            set_rtl(para)

    def line_with_right(self, left, right, bold_left=True):
        """'Role — Employer, City' ........ 'dates' on one line (tab stop, no tables → ATS-safe)."""
        para = self.doc.add_paragraph()
        sec = self.doc.sections[0]
        width = sec.page_width - sec.left_margin - sec.right_margin
        from docx.enum.text import WD_TAB_ALIGNMENT
        para.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
        r = para.add_run(left)
        r.bold = bold_left
        if right:
            para.add_run("\t" + right)
        para.paragraph_format.space_before = Pt(4)
        para.paragraph_format.space_after = Pt(1)
        if self.rtl:
            set_rtl(para)
        return para

    def build(self, out):
        d = self.d
        center = WD_ALIGN_PARAGRAPH.CENTER
        if d.get("photo"):
            try:
                self.doc.add_picture(d["photo"], width=Cm(3))
                self.doc.paragraphs[-1].alignment = center
            except Exception as e:  # photo is optional; never fail the CV for it
                print(f"warning: photo skipped ({e})", file=sys.stderr)
        self.p(d.get("name", ""), size=20, bold=True, color=ACCENT, align=center, after=0)
        if d.get("title"):
            self.p(d["title"], size=12, color=RGBColor(0x4F, 0x46, 0xE5), align=center, after=2)
        c = d.get("contact") or {}
        place = ", ".join(x for x in (c.get("city"), c.get("country")) if x)
        bits = [x for x in (place, c.get("phone"), c.get("email"), *(c.get("links") or [])) if x]
        if bits:
            self.p("  |  ".join(bits), size=9.5, align=center, after=6)

        if d.get("summary"):
            self.heading("summary")
            self.p(d["summary"])

        exp, edu = d.get("experience") or [], d.get("education") or []
        edu_first = d.get("education_first")
        if edu_first is None:
            edu_first = not exp
        order = ["education", "experience"] if edu_first else ["experience", "education"]
        for key in order:
            if key == "experience" and exp:
                self.heading("experience")
                for e in exp:
                    left = " — ".join(x for x in (e.get("role"), ", ".join(y for y in (e.get("employer"), e.get("city")) if y)) if x)
                    dates = " – ".join(x for x in (e.get("start"), e.get("end")) if x)
                    self.line_with_right(left, dates)
                    for b in e.get("bullets") or []:
                        self.bullet(b)
            if key == "education" and edu:
                self.heading("education")
                for e in edu:
                    left = " — ".join(x for x in (e.get("degree"), ", ".join(y for y in (e.get("school"), e.get("city")) if y)) if x)
                    self.line_with_right(left, str(e.get("year") or ""))
                    if e.get("details"):
                        self.p(e["details"], size=10, italic=True)

        if d.get("skills"):
            self.heading("skills")
            for g in d["skills"]:
                items = ", ".join(g.get("items") or [])
                para = self.doc.add_paragraph()
                if g.get("group"):
                    para.add_run(g["group"] + ": ").bold = True
                para.add_run(items)
                para.paragraph_format.space_after = Pt(1)
                if self.rtl:
                    set_rtl(para)
        if d.get("languages"):
            self.heading("languages")
            self.p("  ·  ".join(f"{x.get('name','')}" + (f" ({x['level']})" if x.get("level") else "") for x in d["languages"]))
        for key in ("certifications", "projects", "extras"):
            if d.get(key):
                self.heading(key)
                for item in d[key]:
                    self.bullet(item)
        self.doc.save(out)


def main():
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    with open(sys.argv[1], encoding="utf-8") as f:
        data = json.load(f)
    if not data.get("name"):
        sys.exit("error: 'name' is required")
    CV(data).build(sys.argv[2])
    print(sys.argv[2])


if __name__ == "__main__":
    main()
